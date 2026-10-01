"""Pure-unit tests for the presentation theme palettes + programmatic pptx renderer.

These tests need NO database. They MUST live under tests/unit/ (not tests/api/)
so the DB-truncating tests/api/conftest.py never runs on collection.
"""
import io

import pytest
from pptx import Presentation as PptxOpen
from PIL import Image

from app.services.presentation_themes import THEMES, get_theme
from app.services.pptx_renderer import PptxRenderer


# --- Task 4: theme palettes ------------------------------------------------

def test_themes_have_four_presets_with_required_fields():
    assert set(THEMES.keys()) >= {'polymind', 'dark', 'minimal', 'vibrant'}
    for key, t in THEMES.items():
        for field in ('bg', 'surface', 'primary', 'text', 'muted', 'title_font', 'body_font'):
            assert field in t, f"{key} missing {field}"


def test_get_theme_falls_back_to_polymind():
    assert get_theme('does-not-exist') is THEMES['polymind']


# --- Task 5: PptxRenderer --------------------------------------------------

_OUTLINE = {
    "title": "گزارش فصلی", "language": "fa",
    "slides": [
        {"layout": "title", "title": "گزارش فصلی", "bullets": [], "image_prompt": None},
        {"layout": "content", "title": "اهداف",
         "bullets": ["رشد فروش", "ورود به بازار جدید"], "image_prompt": None},
        {"layout": "closing", "title": "سپاسگزاریم", "bullets": [], "image_prompt": None},
    ],
}


def test_renderer_builds_deck_with_slide_per_outline_item():
    data = PptxRenderer().render(_OUTLINE, theme='polymind', language='fa')
    assert len(PptxOpen(io.BytesIO(data)).slides) == 3


def test_renderer_sets_rtl_on_persian_paragraphs():
    data = PptxRenderer().render(_OUTLINE, theme='polymind', language='fa')
    xml = PptxOpen(io.BytesIO(data)).slides[1]._element.xml
    assert 'rtl="1"' in xml


def test_renderer_writes_title_text():
    data = PptxRenderer().render(_OUTLINE, theme='polymind', language='fa')
    prs = PptxOpen(io.BytesIO(data))
    texts = [s.text for sl in prs.slides for s in sl.shapes if s.has_text_frame]
    assert any('اهداف' in tx for tx in texts)


# --- Task 6: image insertion ----------------------------------------------

def test_renderer_inserts_image_when_bytes_provided():
    buf = io.BytesIO()
    Image.new('RGB', (32, 32), (200, 30, 30)).save(buf, format='PNG')
    outline = {"title": "t", "language": "fa",
               "slides": [{"layout": "content_image", "title": "با تصویر",
                           "bullets": ["یک"], "image_prompt": "x"}]}
    data = PptxRenderer().render(outline, theme='polymind', language='fa', images={0: buf.getvalue()})
    prs = PptxOpen(io.BytesIO(data))
    pics = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]  # 13 == PICTURE
    assert len(pics) == 1


def test_renderer_skips_corrupt_image_without_failing_deck():
    """Spec §5.4: a bad image renders the slide without it, never aborts the deck."""
    outline = {"title": "t", "language": "fa",
               "slides": [{"layout": "content_image", "title": "خراب",
                           "bullets": ["یک"], "image_prompt": "x"}]}
    data = PptxRenderer().render(outline, theme='polymind', language='fa',
                                 images={0: b'not-a-real-image'})
    prs = PptxOpen(io.BytesIO(data))
    assert len(prs.slides) == 1  # deck still built
    pics = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]
    assert pics == []  # corrupt image silently skipped
    # title text still present
    texts = [s.text for s in prs.slides[0].shapes if s.has_text_frame]
    assert any('خراب' in tx for tx in texts)


def test_renderer_preserves_image_aspect_ratio():
    """A wide (non-square) image must not be stretched to the square box."""
    buf = io.BytesIO()
    Image.new('RGB', (400, 100), (30, 60, 200)).save(buf, format='PNG')  # 4:1
    outline = {"title": "t", "language": "fa",
               "slides": [{"layout": "content_image", "title": "نسبت",
                           "bullets": ["یک"], "image_prompt": "x"}]}
    data = PptxRenderer().render(outline, theme='polymind', language='fa', images={0: buf.getvalue()})
    prs = PptxOpen(io.BytesIO(data))
    pics = [sh for sh in prs.slides[0].shapes if sh.shape_type == 13]
    assert len(pics) == 1
    pic = pics[0]
    # width should dominate (limiting dim); height ~= width/4, NOT equal to width
    assert pic.width > pic.height
    ratio = pic.width / pic.height
    assert 3.5 < ratio < 4.5, f"aspect not preserved (ratio={ratio})"


def test_renderer_ltr_path_has_no_rtl_marker():
    data = PptxRenderer().render(_OUTLINE, theme='polymind', language='en')
    prs = PptxOpen(io.BytesIO(data))
    assert len(prs.slides) == 3
    xml = prs.slides[1]._element.xml
    assert 'rtl="1"' not in xml  # LTR paragraphs carry no rtl flag (default-left)


def test_renderer_writes_speaker_notes():
    outline = {"title": "t", "language": "fa",
               "slides": [{"layout": "content", "title": "عنوان", "bullets": ["یک"],
                           "image_prompt": None, "speaker_notes": "یادداشت گوینده"}]}
    data = PptxRenderer().render(outline, theme='polymind', language='fa')
    prs = PptxOpen(io.BytesIO(data))
    assert prs.slides[0].notes_slide.notes_text_frame.text == 'یادداشت گوینده'


def test_renderer_section_layout_renders_title():
    outline = {"title": "t", "language": "fa",
               "slides": [{"layout": "section", "title": "بخش", "bullets": [],
                           "image_prompt": None}]}
    data = PptxRenderer().render(outline, theme='polymind', language='fa')
    prs = PptxOpen(io.BytesIO(data))
    assert len(prs.slides) == 1
    texts = [s.text for s in prs.slides[0].shapes if s.has_text_frame]
    assert any('بخش' in tx for tx in texts)


@pytest.mark.parametrize('theme', ['polymind', 'dark', 'minimal', 'vibrant'])
def test_renderer_smoke_over_all_themes(theme):
    data = PptxRenderer().render(_OUTLINE, theme=theme, language='fa')
    prs = PptxOpen(io.BytesIO(data))
    assert len(prs.slides) == 3
