"""Tests for the payroll payslip core (pure functions — no DB, no FastAPI).

These do not use any conftest fixture, so they neither touch Postgres nor the
autouse truncation path.

Self-validation anchors come from the اردیبهشت test row (سپهر رادمرد / 10415).
The values are taken from the *paid* مزایای مستمر/غیرمستمر bands (NOT the حکم
decree columns) and reconcile arithmetically:

    earnings_total(256,652,286) − بیمه(17,965,660) − وام(15,000,000)
        = 223,686,626 = خالص پرداختنی (net)
"""
from __future__ import annotations

import io
import os
import zipfile

import pytest

from app.services.payroll_service import build_payslip_zip, parse_payroll
from app.services.payslip_render import render_payslip

_XLSX_PATH = os.path.join(
    os.path.dirname(__file__), "..", "..", "..",
    "test-data", "حقوق و دستمزد 14050000.xlsx",
)


@pytest.fixture(scope="module")
def xlsx_bytes() -> bytes:
    if not os.path.exists(_XLSX_PATH):
        pytest.skip(f"payroll test file missing: {_XLSX_PATH}")
    with open(_XLSX_PATH, "rb") as f:
        return f.read()


@pytest.fixture(scope="module")
def parsed(xlsx_bytes) -> dict:
    return parse_payroll(xlsx_bytes)


@pytest.fixture(scope="module")
def record(parsed) -> dict:
    return parsed["records"][0]


def _earning(record: dict, label: str) -> int:
    for e in record["earnings"]:
        if e["label"] == label:
            return e["amount"]
    raise AssertionError(f"earning line {label!r} not found in {record['earnings']}")


def _deduction(record: dict, label: str) -> int:
    for d in record["deductions"]:
        if d["label"] == label:
            return d["amount"]
    raise AssertionError(f"deduction line {label!r} not found in {record['deductions']}")


# ---------------------------------------------------------------------------
# parse_payroll
# ---------------------------------------------------------------------------
def test_parse_structure(parsed):
    assert parsed["columns_ok"] is True
    assert parsed["month"] == "اردیبهشت"
    assert parsed["employer"] == "شرکت نمونه"
    assert len(parsed["records"]) == 1
    assert parsed["warnings"] == []


def test_parse_identity(record):
    assert record["code"] == "10415"
    assert record["name"] == "سپهر رادمرد"
    assert record["work_days"] == 31
    assert record["overtime_hours"] == 31.04


def test_parse_earnings_anchors(record):
    # PAID مستمر/غیرمستمر amounts, NOT the حکم decree base.
    assert _earning(record, "حقوق پایه") == 171_797_355
    assert _earning(record, "حق مسکن") == 30_000_000
    assert _earning(record, "حق خواروبار") == 22_000_000  # renamed from حق بن
    assert record["earnings_total"] == 256_652_286


def test_parse_earnings_total_is_authoritative(record):
    # AL (جمع کل مستمر) + AW (جمع کل غیرمستمر), authoritative.
    line_sum = sum(e["amount"] for e in record["earnings"])
    assert abs(line_sum - record["earnings_total"]) <= 1


def test_parse_deductions_anchors(record):
    assert _deduction(record, "بیمه تامین اجتماعی سهم کارمند") == 17_965_660
    assert _deduction(record, "مالیات") == 0  # shown even though zero
    assert record["deductions_total"] == 17_965_660


def test_parse_installment_advance_net(record):
    assert record["installment"] == 15_000_000  # وام
    assert record["advance"] == 0               # مساعده
    assert record["net"] == 223_686_626         # خالص پرداختنی, used directly


def test_net_reconciles_with_earnings_and_deductions(record):
    # earnings_total − deductions_total − installment == net (internal proof).
    derived = record["earnings_total"] - record["deductions_total"] - record["installment"]
    assert derived == record["net"]


def test_earnings_have_no_zero_lines_except_always_show(record):
    for e in record["earnings"]:
        assert e["amount"] != 0 or e["label"] == "حقوق پایه"


def test_parse_missing_columns_does_not_crash():
    # A workbook with the wrong shape → columns_ok False + warnings, no raise.
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "تست"
    ws["A1"] = "چیزی"
    ws["A2"] = "نامربوط"
    buf = io.BytesIO()
    wb.save(buf)

    res = parse_payroll(buf.getvalue())
    assert res["columns_ok"] is False
    assert res["warnings"]
    assert res["records"] == []
    assert res["month"] == "تست"


def test_employer_and_month_overrides(xlsx_bytes):
    res = parse_payroll(xlsx_bytes, employer="شرکت آزمون", month_override="خرداد")
    assert res["employer"] == "شرکت آزمون"
    assert res["month"] == "خرداد"


# ---------------------------------------------------------------------------
# render_payslip
# ---------------------------------------------------------------------------
def test_render_returns_pdf_bytes(record, parsed):
    pdf = render_payslip(record, employer=parsed["employer"], month=parsed["month"])
    assert isinstance(pdf, bytes)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 1500  # a real rendered page, not an empty stub


def test_render_is_idempotent_on_fonts(record, parsed):
    # Two renders in a row must both succeed (font re-registration guard).
    a = render_payslip(record, employer=parsed["employer"], month=parsed["month"])
    b = render_payslip(record, employer=parsed["employer"], month=parsed["month"])
    assert a.startswith(b"%PDF") and b.startswith(b"%PDF")


# ---------------------------------------------------------------------------
# build_payslip_zip
# ---------------------------------------------------------------------------
def test_build_zip_25_employees(record, parsed):
    synthetic = []
    for i in range(25):
        clone = dict(record)
        clone["name"] = f"کارمند شماره {i + 1}"
        clone["code"] = f"{10000 + i}"
        synthetic.append(clone)

    seen: list[tuple[int, int]] = []
    zip_bytes = build_payslip_zip(
        synthetic,
        employer=parsed["employer"],
        month=parsed["month"],
        progress_cb=lambda done, total: seen.append((done, total)),
    )

    assert zip_bytes[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        names = zf.namelist()
        assert len(names) == 25
        assert all(n.endswith(".pdf") for n in names)
        assert all(zf.read(n).startswith(b"%PDF") for n in names)

    # progress_cb fired once per slip, monotonically, ending at (25, 25).
    assert seen == [(i, 25) for i in range(1, 26)]


def test_build_zip_dedupes_identical_names(record, parsed):
    dupes = [dict(record), dict(record), dict(record)]
    zip_bytes = build_payslip_zip(dupes, employer=parsed["employer"], month=parsed["month"])
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        names = zf.namelist()
        assert len(names) == 3
        assert len(set(names)) == 3  # no silent overwrite


def test_zip_entry_name_preserves_persian(record, parsed):
    zip_bytes = build_payslip_zip([record], employer=parsed["employer"], month=parsed["month"])
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        name = zf.namelist()[0]
        # Persian glyphs preserved; the inner space is collapsed to '_' (the
        # unicode-preserving sanitizer joins on whitespace, like display_filename).
        assert "سپهر" in name and "رادمرد" in name
        assert name.endswith("10415.pdf")
