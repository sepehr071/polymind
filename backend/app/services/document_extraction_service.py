"""Document text extraction for chat attachments.

Converts uploaded office / text documents to plain text so their content can be
threaded into chat prompts. Native PDFs are NOT extracted here — they are
forwarded to the model as a ``data:application/pdf;base64,...`` data URL (PDF
OCR happens model-side via OpenRouter's file-parser), so this module only
exposes the helpers the upload route needs to branch.

Extraction uses lightweight, pure-Python libraries (no ML/onnxruntime):
  * ``.docx`` → python-docx
  * ``.doc``  → olefile (crude OLE string recovery) + optional antiword/catdoc
  * ``.xlsx`` → openpyxl       * ``.xls`` → xlrd
  * ``.pptx`` → python-pptx
  * ``.csv`` / ``.txt`` / ``.md`` / ``.json`` / ``.xml`` → stdlib decode
  * ``.html`` / ``.htm`` → stdlib decode + tag strip

Every library import is GUARDED: a missing lib degrades that format to
``status='unavailable'`` rather than crashing the upload path. Nothing here
raises — ``extract_text`` always returns the contract dict
``{'status', 'markdown', 'chars', 'truncated'}``.
"""
from __future__ import annotations

import base64
import csv as _csv
import html as _html
import logging
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from io import BytesIO, StringIO

logger = logging.getLogger(__name__)

# Extensions we can turn into text. PDFs are handled separately (native data-URL
# passthrough), so they are NOT in this set.
EXTRACTABLE_EXTS = {
    'doc', 'docx', 'xlsx', 'xls', 'pptx', 'csv', 'html', 'htm', 'txt', 'md',
    'json', 'xml',
}
NATIVE_PDF_EXTS = {'pdf'}
# Archive formats unpacked + recursively text-extracted in-process (stdlib only).
ZIP_EXTS = {'zip'}
# Raster image extensions surfaced out of a ZIP (forwarded to the model as data
# URLs by the caller). Maps each to its MIME type.
_ZIP_IMAGE_MIMES = {
    'png': 'image/png',
    'jpg': 'image/jpeg',
    'jpeg': 'image/jpeg',
    'webp': 'image/webp',
    'gif': 'image/gif',
}

# Office formats (.docx/.xlsx/.pptx) ARE zip archives, so a directly-uploaded
# one is a zip-bomb vector even before its parser library sees it (the OUTPUT
# char cap clips text only AFTER openpyxl/python-docx/python-pptx has already
# inflated everything in memory). These pre-parse caps mirror the ZIP-attachment
# limits below: a raw byte ceiling, a total-uncompressed ceiling, a
# compression-ratio ceiling, and a member-count ceiling.
_OFFICE_XML_EXTS = {'docx', 'xlsx', 'pptx'}
_OFFICE_MAX_RAW_BYTES = 50 * 1024 * 1024       # reject the upload itself past this
# Σ member uncompressed sizes. Kept ~3x the 32 MB ingress cap: comfortably
# covers any genuine office document while keeping the worst-case transient
# inflation (.docx/.pptx do a full non-streaming lxml DOM parse) well bounded.
_OFFICE_MAX_UNCOMPRESSED_BYTES = 96 * 1024 * 1024  # was 200 MB (tracks ingress)
_OFFICE_MAX_COMPRESSION_RATIO = 80             # total uncompressed / compressed
_OFFICE_MAX_MEMBERS = 5000                     # absurd member count = crafted bomb

# Guarded per-format imports — a missing lib degrades that format to
# 'unavailable' instead of taking down the module at import time.
try:  # python-docx
    import docx as _docx  # type: ignore
    _DOCX_OK = True
except Exception:  # noqa: BLE001
    _docx = None  # type: ignore
    _DOCX_OK = False

try:  # olefile — legacy binary .doc (OLE Compound File)
    import olefile as _olefile  # type: ignore
    _OLE_OK = True
except Exception:  # noqa: BLE001
    _olefile = None  # type: ignore
    _OLE_OK = False

# Optional external tools for higher-fidelity .doc text (used when on PATH).
_ANTIWORD_BIN = shutil.which('antiword')
_CATDOC_BIN = shutil.which('catdoc')
# .doc is extractable when olefile is present OR an external tool is available.
_DOC_OK = _OLE_OK or bool(_ANTIWORD_BIN or _CATDOC_BIN)

try:  # openpyxl
    from openpyxl import load_workbook as _load_xlsx  # type: ignore
    _XLSX_OK = True
except Exception:  # noqa: BLE001
    _load_xlsx = None  # type: ignore
    _XLSX_OK = False

try:  # xlrd (legacy .xls)
    import xlrd as _xlrd  # type: ignore
    _XLS_OK = True
except Exception:  # noqa: BLE001
    _xlrd = None  # type: ignore
    _XLS_OK = False

try:  # python-pptx
    from pptx import Presentation as _Presentation  # type: ignore
    _PPTX_OK = True
except Exception:  # noqa: BLE001
    _Presentation = None  # type: ignore
    _PPTX_OK = False


def is_extractable(ext: str) -> bool:
    """True iff ``ext`` (with or without a leading dot) can be text-extracted."""
    return (ext or '').lstrip('.').lower() in EXTRACTABLE_EXTS


def is_native_pdf(ext: str) -> bool:
    """True iff ``ext`` is a PDF (forwarded as a data URL, not text-extracted)."""
    return (ext or '').lstrip('.').lower() in NATIVE_PDF_EXTS


def is_zip(ext: str) -> bool:
    """True iff ``ext`` is a ZIP archive (unpacked + recursively extracted)."""
    return (ext or '').lstrip('.').lower() in ZIP_EXTS


def pdf_data_url(file_bytes: bytes) -> str:
    """Base64-encode ``file_bytes`` as an OpenRouter PDF data URL."""
    encoded = base64.b64encode(file_bytes or b'').decode('ascii')
    return f"data:application/pdf;base64,{encoded}"


# --------------------------------------------------------------------------
# Per-format extractors. Each raises on failure; the public wrapper catches.
# --------------------------------------------------------------------------
def _extract_docx(file_bytes: bytes) -> str:
    document = _docx.Document(BytesIO(file_bytes))
    parts: list[str] = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


def _doc_via_external_tool(file_bytes: bytes) -> str | None:
    """Higher-fidelity .doc extract via antiword/catdoc when installed.

    Returns plain text on success, ``None`` when no tool is available or the
    subprocess fails. Never raises — caller falls back to olefile recovery.
    """
    bin_path = _ANTIWORD_BIN or _CATDOC_BIN
    if not bin_path or not file_bytes:
        return None
    tmp_path = None
    try:
        fd, tmp_path = tempfile.mkstemp(suffix='.doc')
        os.close(fd)
        with open(tmp_path, 'wb') as fh:
            fh.write(file_bytes)
        # antiword: stdout is UTF-8-ish text; catdoc: same with -w (word wrap off).
        if bin_path == _ANTIWORD_BIN:
            argv = [bin_path, '-m', 'UTF-8.txt', tmp_path]
        else:
            argv = [bin_path, '-w', tmp_path]
        proc = subprocess.run(
            argv,
            capture_output=True,
            timeout=20,
            check=False,
        )
        if proc.returncode != 0:
            return None
        out = (proc.stdout or b'')
        if isinstance(out, bytes):
            text = out.decode('utf-8', errors='replace')
        else:
            text = str(out)
        text = text.strip()
        return text or None
    except Exception:  # noqa: BLE001 — external tool is best-effort
        logger.debug("doc external tool extract failed", exc_info=True)
        return None
    finally:
        if tmp_path:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


def _recover_printable_runs(raw: bytes) -> str:
    """Pull readable UTF-16LE + ASCII runs out of a Word OLE stream.

    Not a full Word 97 parser — good enough for DLP / prompt injection of plain
    prose, emails, codenames, and table cell strings embedded as text runs.
    """
    parts: list[str] = []
    # UTF-16LE runs: ≥4 consecutive printable BMP chars (8+ raw bytes).
    for m in re.finditer(
        rb'(?:[\x20-\x7e\xa0-\xff]\x00){4,}',
        raw,
    ):
        try:
            s = m.group().decode('utf-16-le', errors='ignore').strip()
        except Exception:  # noqa: BLE001
            continue
        if s and not s.isspace():
            parts.append(s)
    # ASCII runs (length ≥ 4) that weren't already covered as UTF-16LE pairs.
    for m in re.finditer(rb'[\x20-\x7e]{4,}', raw):
        s = m.group().decode('ascii', errors='ignore').strip()
        if s and not s.isspace():
            parts.append(s)
    # De-dupe while preserving order (UTF-16 + ASCII often double-hit).
    seen: set[str] = set()
    out: list[str] = []
    for p in parts:
        if p in seen:
            continue
        seen.add(p)
        out.append(p)
    return '\n'.join(out)


def _extract_doc(file_bytes: bytes) -> str:
    """Legacy binary Word (``.doc`` / Word 97–2003 OLE) → plain text.

    Prefers antiword/catdoc when present (higher fidelity). Falls back to
    olefile stream recovery so DLP still sees plain-text secrets without a
    system package.
    """
    external = _doc_via_external_tool(file_bytes)
    if external:
        return external

    if not _OLE_OK or _olefile is None:
        raise RuntimeError('olefile unavailable and no antiword/catdoc on PATH')

    if not _olefile.isOleFile(BytesIO(file_bytes)):
        # Some "doc" uploads are mislabeled RTF/plain — try a soft decode.
        try:
            as_text = file_bytes.decode('utf-8', errors='replace').strip()
        except Exception:  # noqa: BLE001
            as_text = ''
        if as_text and sum(1 for c in as_text if c.isprintable() or c in '\r\n\t') > len(as_text) * 0.7:
            return as_text
        raise ValueError('not a valid OLE compound .doc file')

    ole = _olefile.OleFileIO(BytesIO(file_bytes))
    try:
        chunks: list[str] = []
        # WordDocument is the primary body stream; 1Table/0Table hold piece
        # tables / extra text. Read whatever exists and recover printable runs.
        for stream_name in (
            'WordDocument',
            '1Table',
            '0Table',
            'Data',
        ):
            if not ole.exists(stream_name):
                continue
            try:
                raw = ole.openstream(stream_name).read()
            except Exception:  # noqa: BLE001
                continue
            recovered = _recover_printable_runs(raw)
            if recovered:
                chunks.append(recovered)
        text = '\n'.join(chunks).strip()
        if not text:
            raise ValueError('no recoverable text in .doc streams')
        return text
    finally:
        try:
            ole.close()
        except Exception:  # noqa: BLE001
            pass


def _extract_xlsx(file_bytes: bytes) -> str:
    wb = _load_xlsx(BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        out: list[str] = []
        for ws in wb.worksheets:
            out.append(f"# {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = ['' if v is None else str(v) for v in row]
                if any(cells):
                    out.append('\t'.join(cells))
        return '\n'.join(out)
    finally:
        wb.close()


def _extract_xls(file_bytes: bytes) -> str:
    book = _xlrd.open_workbook(file_contents=file_bytes)
    out: list[str] = []
    for sheet in book.sheets():
        out.append(f"# {sheet.name}")
        for rx in range(sheet.nrows):
            cells = [str(sheet.cell_value(rx, cx)) for cx in range(sheet.ncols)]
            if any(c.strip() for c in cells):
                out.append('\t'.join(cells))
    return '\n'.join(out)


def _extract_pptx(file_bytes: bytes) -> str:
    prs = _Presentation(BytesIO(file_bytes))
    out: list[str] = []
    for idx, slide in enumerate(prs.slides, 1):
        out.append(f"## Slide {idx}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = ''.join(run.text for run in para.runs).strip()
                    if line:
                        out.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        out.append(' | '.join(cells))
    return '\n'.join(out)


def _extract_csv(file_bytes: bytes) -> str:
    text = file_bytes.decode('utf-8', errors='replace')
    rows = list(_csv.reader(StringIO(text)))
    return '\n'.join('\t'.join(r) for r in rows)


def _strip_html(file_bytes: bytes) -> str:
    raw = file_bytes.decode('utf-8', errors='replace')
    raw = re.sub(r'<(script|style)\b[^>]*>.*?</\1>', ' ', raw, flags=re.S | re.I)
    raw = re.sub(r'<[^>]+>', ' ', raw)
    raw = _html.unescape(raw)
    return re.sub(r'[ \t]*\n[ \t]*(\n[ \t]*)+', '\n\n', re.sub(r'[ \t]+', ' ', raw)).strip()


def _extract_plaintext(file_bytes: bytes) -> str:
    return file_bytes.decode('utf-8', errors='replace')


class _ZipBombError(Exception):
    """Raised when a directly-uploaded office file trips a zip-bomb guard."""


def _guard_office_xml(file_bytes: bytes) -> None:
    """Pre-parse zip-bomb guard for a directly-uploaded .docx/.xlsx/.pptx.

    Office files are zip containers; this inspects the central directory WITHOUT
    decompressing any member and raises :class:`_ZipBombError` when the raw size,
    total uncompressed size, compression ratio, or member count exceeds a sane
    bound. The public :func:`extract_text` wrapper maps that to ``status='error'``
    so a crafted bomb never reaches openpyxl/python-docx/python-pptx.
    """
    if len(file_bytes) > _OFFICE_MAX_RAW_BYTES:
        raise _ZipBombError(
            f"file too large ({len(file_bytes)} bytes > {_OFFICE_MAX_RAW_BYTES})"
        )
    # A non-zip payload (e.g. legacy binary mislabeled .docx) raises BadZipFile;
    # let it propagate to the wrapper's generic handler as a normal error.
    with zipfile.ZipFile(BytesIO(file_bytes)) as zf:
        infos = zf.infolist()
        if len(infos) > _OFFICE_MAX_MEMBERS:
            raise _ZipBombError(
                f"too many members ({len(infos)} > {_OFFICE_MAX_MEMBERS})"
            )
        total_uncompressed = 0
        total_compressed = 0
        for info in infos:
            total_uncompressed += info.file_size
            total_compressed += info.compress_size
            if total_uncompressed > _OFFICE_MAX_UNCOMPRESSED_BYTES:
                raise _ZipBombError(
                    f"uncompressed size exceeds {_OFFICE_MAX_UNCOMPRESSED_BYTES} bytes"
                )
        if (
            total_compressed > 0
            and (total_uncompressed / total_compressed) > _OFFICE_MAX_COMPRESSION_RATIO
        ):
            raise _ZipBombError(
                f"compression ratio {total_uncompressed / total_compressed:.0f}x "
                f"exceeds {_OFFICE_MAX_COMPRESSION_RATIO}x"
            )


# ext -> (extractor, availability-flag). None flag = always available (stdlib).
_EXTRACTORS = {
    'doc': (_extract_doc, '_DOC_OK'),
    'docx': (_extract_docx, '_DOCX_OK'),
    'xlsx': (_extract_xlsx, '_XLSX_OK'),
    'xls': (_extract_xls, '_XLS_OK'),
    'pptx': (_extract_pptx, '_PPTX_OK'),
    'csv': (_extract_csv, None),
    'html': (_strip_html, None),
    'htm': (_strip_html, None),
    'txt': (_extract_plaintext, None),
    'md': (_extract_plaintext, None),
    'json': (_extract_plaintext, None),
    'xml': (_extract_plaintext, None),
}


def extract_text(
    file_bytes: bytes,
    *,
    filename: str,
    mime_type: str,
    extension: str,
    max_chars: int,
) -> dict:
    """Extract plain text from ``file_bytes`` based on ``extension``.

    Returns ``{'status', 'markdown', 'chars', 'truncated'}``:
      * ``status='unavailable'`` — no extractor / its library is missing.
      * ``status='error'``       — extraction raised (logged; empty text).
      * ``status='ok'``          — extracted, under ``max_chars``.
      * ``status='truncated'``   — extracted then clipped to ``max_chars`` with
        a trailing ``[...truncated N of M characters...]`` marker.

    Never raises — every failure mode maps to a contract dict.
    """
    ext = (extension or '').lstrip('.').lower()
    entry = _EXTRACTORS.get(ext)
    if entry is None:
        return {'status': 'unavailable', 'markdown': '', 'chars': 0, 'truncated': False}

    extractor, flag = entry
    if flag is not None and not globals().get(flag, False):
        return {'status': 'unavailable', 'markdown': '', 'chars': 0, 'truncated': False}

    # Office formats are zip containers — guard against a zip bomb BEFORE the
    # parser library inflates the archive in memory.
    if ext in _OFFICE_XML_EXTS:
        try:
            _guard_office_xml(file_bytes)
        except _ZipBombError as exc:
            logger.warning("office zip-bomb guard rejected %s: %s", filename, exc)
            return {'status': 'error', 'markdown': '', 'chars': 0, 'truncated': False}

    try:
        text = extractor(file_bytes) or ''
    except Exception as exc:  # noqa: BLE001 - extraction must never crash upload
        logger.exception("document extraction failed for %s: %s", filename, exc)
        return {'status': 'error', 'markdown': '', 'chars': 0, 'truncated': False}

    truncated = False
    if len(text) > max_chars:
        dropped = len(text) - max_chars
        original_len = len(text)
        text = text[:max_chars] + (
            f"\n\n[...truncated {dropped} of {original_len} characters...]"
        )
        truncated = True

    return {
        'status': 'truncated' if truncated else 'ok',
        'markdown': text,
        'chars': len(text),
        'truncated': truncated,
    }


# --------------------------------------------------------------------------
# ZIP archive extraction. Unpacks an archive in-process (stdlib zipfile only)
# and recursively text-extracts inner files via ``extract_text``. Hardened
# against zip-bombs and path-traversal; never raises.
# --------------------------------------------------------------------------
_ZIP_PER_FILE_MAX_BYTES = 25 * 1024 * 1024     # skip any single entry larger than this
_ZIP_TOTAL_BUDGET_BYTES = 100 * 1024 * 1024    # cumulative uncompressed read budget
_ZIP_MAX_IMAGES = 10                           # max raster images surfaced
_ZIP_IMAGE_MAX_BYTES = 8 * 1024 * 1024         # skip any image larger than this


def _zip_ext(name: str) -> str:
    """Lowercase extension (no dot) of an archive member path, '' if none."""
    base = name.rsplit('/', 1)[-1]
    return base.rsplit('.', 1)[-1].lower() if '.' in base else ''


def _zip_human_size(num_bytes: int) -> str:
    """Compact human-readable size for the file-tree header (e.g. '12 KB')."""
    size = float(num_bytes)
    for unit in ('B', 'KB', 'MB', 'GB'):
        if size < 1024 or unit == 'GB':
            if unit == 'B':
                return f"{int(size)} {unit}"
            return f"{size:.0f} {unit}"
        size /= 1024
    return f"{int(num_bytes)} B"


def _zip_is_symlink(info: zipfile.ZipInfo) -> bool:
    """True iff the entry's stored Unix mode marks it as a symlink."""
    return ((info.external_attr >> 16) & 0o170000) == 0o120000


def _zip_should_skip(info: zipfile.ZipInfo) -> bool:
    """Skip directories, macOS resource forks, dotfiles, traversal + symlinks."""
    name = info.filename
    if info.is_dir() or name.endswith('/'):
        return True
    if name.startswith('__MACOSX/'):
        return True
    basename = name.rsplit('/', 1)[-1]
    if basename.startswith('.'):
        return True
    # Path traversal / absolute paths: never read these.
    if name.startswith('/') or name.startswith('\\') or '..' in name.replace('\\', '/').split('/'):
        return True
    if _zip_is_symlink(info):
        return True
    return False


def extract_zip(file_bytes: bytes, *, max_chars: int, max_entries: int = 200) -> dict:
    """Unpack a ZIP archive and recursively text-extract its inner files.

    Returns::

        {
          'status': 'ok'|'truncated'|'error',
          'markdown': str,   # file-tree header + per-file '## <path>' bodies
          'chars': int,
          'truncated': bool,
          'entries': [ {'name', 'size', 'kind': 'text'|'image'|'other'} ],
          'images': [ {'name', 'bytes_b64', 'mime'} ],
        }

    Hardened against zip-bombs (per-file + total uncompressed caps, entry cap)
    and path-traversal/symlink entries. Never raises — any failure (including a
    corrupt archive) maps to ``status='error'`` with empty payloads.
    """
    empty_error = {
        'status': 'error', 'markdown': '', 'chars': 0,
        'truncated': False, 'entries': [], 'images': [],
    }

    try:
        zf = zipfile.ZipFile(BytesIO(file_bytes))
    except Exception as exc:  # noqa: BLE001 - bad/corrupt archive must not crash upload
        logger.warning("zip open failed: %s", exc)
        return empty_error

    truncated = False
    entries: list[dict] = []
    images: list[dict] = []
    # (info, ext) pairs for text/image members we may read, in archive order.
    text_members: list[tuple[zipfile.ZipInfo, str]] = []
    image_members: list[tuple[zipfile.ZipInfo, str]] = []

    try:
        processed = 0
        for info in zf.infolist():
            if _zip_should_skip(info):
                continue
            if processed >= max_entries:
                # Hit the entry cap: more files exist than we will list/read.
                truncated = True
                break
            processed += 1

            name = info.filename
            ext = _zip_ext(name)
            if is_extractable(ext):
                kind = 'text'
                text_members.append((info, ext))
            elif ext in _ZIP_IMAGE_MIMES:
                kind = 'image'
                image_members.append((info, ext))
            else:
                kind = 'other'
            entries.append({'name': name, 'size': info.file_size, 'kind': kind})

        # ---- File-tree header (lists every entry we catalogued) ------------
        header_lines = [f"ZIP contents ({len(entries)} files):"]
        for e in entries:
            header_lines.append(
                f"- {e['name']} ({e['kind']}, {_zip_human_size(e['size'])})"
            )
        markdown = '\n'.join(header_lines)

        total_read = 0  # cumulative uncompressed bytes read (zip-bomb budget)

        def _read_member(info: zipfile.ZipInfo, byte_cap: int) -> bytes | None:
            """Read a member's bytes, honoring per-file + total budgets.

            Returns None (and marks the archive truncated) if the entry is too
            large or the total budget is exhausted.
            """
            nonlocal total_read, truncated
            if info.file_size > byte_cap:
                truncated = True
                return None
            if total_read + info.file_size > _ZIP_TOTAL_BUDGET_BYTES:
                truncated = True
                return None
            try:
                data = zf.read(info)
            except Exception as exc:  # noqa: BLE001 - skip unreadable member, keep going
                logger.warning("zip member read failed for %s: %s", info.filename, exc)
                return None
            total_read += len(data)
            return data

        # ---- Text bodies (name order), tracked against max_chars ----------
        char_budget = max_chars - len(markdown)
        for info, ext in sorted(text_members, key=lambda m: m[0].filename):
            if char_budget <= 0:
                truncated = True
                break
            data = _read_member(info, _ZIP_PER_FILE_MAX_BYTES)
            if data is None:
                continue
            inner = extract_text(
                data,
                filename=info.filename,
                mime_type='',
                extension=ext,
                max_chars=char_budget,
            )
            body = inner.get('markdown') or ''
            if not body:
                continue
            section = f"\n\n## {info.filename}\n{body}"
            markdown += section
            char_budget -= len(section)
            if inner.get('truncated'):
                truncated = True

        # ---- Inner raster images (capped count + per-image size) ----------
        for info, ext in image_members:
            if len(images) >= _ZIP_MAX_IMAGES:
                truncated = True
                break
            data = _read_member(info, min(_ZIP_IMAGE_MAX_BYTES, _ZIP_PER_FILE_MAX_BYTES))
            if data is None:
                continue
            images.append({
                'name': info.filename,
                'bytes_b64': base64.b64encode(data).decode('ascii'),
                'mime': _ZIP_IMAGE_MIMES[ext],
            })
    except Exception as exc:  # noqa: BLE001 - any unexpected failure → error contract
        logger.exception("zip extraction failed: %s", exc)
        return empty_error
    finally:
        try:
            zf.close()
        except Exception:  # noqa: BLE001
            pass

    return {
        'status': 'truncated' if truncated else 'ok',
        'markdown': markdown,
        'chars': len(markdown),
        'truncated': truncated,
        'entries': entries,
        'images': images,
    }
