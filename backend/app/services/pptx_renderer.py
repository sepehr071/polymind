"""Programmatic editable .pptx renderer for presentation decks.

Layout catalog is dispatched by slide['layout']. Shared chrome (bg bar, footer,
cards, KPI tiles) keeps themes consistent. RTL uses OOXML a:pPr/@rtl only —
never pre-reshape Persian text.
"""

from __future__ import annotations

import io
from typing import Optional, Protocol

from PIL import Image
from pptx import Presentation as PptxNew
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.services.presentation_themes import get_theme

_RTL_LANGS = frozenset({'fa', 'ar', 'he', 'ur'})
_W = Inches(13.333)
_H = Inches(7.5)
_MARGIN = Inches(0.55)
_BODY_TOP = Inches(1.55)
_IMG_BOX = Inches(4.5)
_FOOTER_H = Inches(0.35)


class DeckRenderer(Protocol):
    def render(self, outline: dict, *, theme: str, language: str,
               images: Optional[dict] = None) -> bytes: ...


def _rgb(t):
    return RGBColor(int(t[0]), int(t[1]), int(t[2]))


def _set_rtl(paragraph, rtl: bool):
    """Mark paragraph RTL via OOXML a:pPr/@rtl. PowerPoint runs BiDi itself;
    do NOT pre-reshape text (that breaks editing)."""
    if not rtl:
        return
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set('rtl', '1')


def _align(rtl: bool):
    return PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT


class PptxRenderer:
    def render(self, outline, *, theme='polymind', language='fa', images=None) -> bytes:
        images = images or {}
        th = get_theme(theme)
        rtl = (language or '').lower() in _RTL_LANGS
        prs = PptxNew()
        prs.slide_width = _W
        prs.slide_height = _H
        blank = prs.slide_layouts[6]
        slides = outline.get('slides') or []
        deck_title = (outline.get('title') or '').strip()
        n = len(slides)
        for idx, slide in enumerate(slides):
            s = prs.slides.add_slide(blank)
            self._bg(s, th)
            layout = (slide.get('layout') or 'content').strip().lower()
            img = images.get(idx)
            # Soft visual_recipe: composition may nudge layout handling
            recipe = slide.get('visual_recipe') if isinstance(slide.get('visual_recipe'), dict) else {}
            if recipe.get('composition') == 'big_number' and layout == 'metrics':
                pass  # metrics handler already card-based
            if recipe.get('composition') == 'full_bleed' and layout in ('title', 'image_hero'):
                layout = 'image_hero' if img else layout
            handler = getattr(self, f'_lay_{layout}', None)
            if handler is None:
                handler = self._lay_content
            handler(s, th, slide, rtl, img)
            if th.get('footer') and layout not in ('title', 'closing', 'section', 'image_hero'):
                self._footer(s, th, deck_title, idx + 1, n, rtl)
            notes = slide.get('speaker_notes')
            if notes:
                s.notes_slide.notes_text_frame.text = notes
        out = io.BytesIO()
        prs.save(out)
        return out.getvalue()

    # ── chrome ──────────────────────────────────────────────────────────

    def _bg(self, s, th):
        rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, _W, _H)
        rect.fill.solid()
        rect.fill.fore_color.rgb = _rgb(th['bg'])
        rect.line.fill.background()
        rect.shadow.inherit = False
        sp = rect._element
        sp.getparent().remove(sp)
        s.shapes._spTree.insert(2, sp)

        style = th.get('bar_style') or 'edge'
        if style == 'edge':
            bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.16), _H)
            bar.fill.solid()
            bar.fill.fore_color.rgb = _rgb(th['primary'])
            bar.line.fill.background()
            bar.shadow.inherit = False
        elif style == 'top':
            bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, _W, Inches(0.12))
            bar.fill.solid()
            bar.fill.fore_color.rgb = _rgb(th['primary'])
            bar.line.fill.background()
            bar.shadow.inherit = False

    def _footer(self, s, th, deck_title, page, total, rtl):
        y = _H - _FOOTER_H - Inches(0.12)
        line = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, _MARGIN, y - Inches(0.08),
            _W - 2 * _MARGIN, Emu(12700),
        )
        line.fill.solid()
        line.fill.fore_color.rgb = _rgb(th['muted'])
        line.line.fill.background()
        line.shadow.inherit = False

        # Page counter always LTR digits so "3 / 12" never reverses under RTL.
        page_txt = f"{int(page)} / {int(total)}"
        title_txt = (deck_title or '')[:60]
        # Physical: title near start edge, page near end edge.
        if rtl:
            title_x, page_x = Inches(2.9), _MARGIN
            title_align, page_align = PP_ALIGN.RIGHT, PP_ALIGN.LEFT
        else:
            title_x, page_x = _MARGIN, _W - _MARGIN - Inches(2.2)
            title_align, page_align = PP_ALIGN.LEFT, PP_ALIGN.RIGHT
        self._textbox(
            s, title_x, y, Inches(8.5), _FOOTER_H,
            title_txt, size=10, color=th['muted'], bold=False, rtl=rtl,
            align=title_align, font=th['body_font'],
        )
        self._textbox(
            s, page_x, y, Inches(2.2), _FOOTER_H,
            page_txt, size=10, color=th['muted'], bold=False, rtl=False,
            align=page_align, font=th['body_font'],
        )

    def _card(self, s, th, x, y, w, h, *, accent_edge=True):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(th['surface'])
        shape.line.fill.background()
        shape.shadow.inherit = False
        if accent_edge:
            edge_w = Inches(0.08)
            edge = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, edge_w, h)
            edge.fill.solid()
            edge.fill.fore_color.rgb = _rgb(th['primary'])
            edge.line.fill.background()
            edge.shadow.inherit = False
        return shape

    def _textbox(self, s, x, y, w, h, text, *, size, color, bold, rtl, align, font,
                 vertical=None):
        box = s.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        if vertical is not None:
            tf.vertical_anchor = vertical
        self._run(tf.paragraphs[0], text or '', size=size, color=color, bold=bold,
                  rtl=rtl, align=align, font=font)
        return box

    def _run(self, p, text, *, size, color, bold, rtl, align, font):
        p.alignment = align
        _set_rtl(p, rtl)
        # clear default empty run if present
        if p.runs:
            run = p.runs[0]
            run.text = text or ''
        else:
            run = p.add_run()
            run.text = text or ''
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)
        run.font.name = font
        return run

    def _title_block(self, s, th, title, rtl, *, kicker=None, subtitle=None):
        align = _align(rtl)
        y = Inches(0.4)
        if kicker:
            self._textbox(
                s, _MARGIN, y, _W - 2 * _MARGIN, Inches(0.35),
                kicker, size=12, color=th['accent'], bold=True, rtl=rtl,
                align=align, font=th['body_font'],
            )
            y = Inches(0.7)
        self._textbox(
            s, _MARGIN, y, _W - 2 * _MARGIN, Inches(0.75),
            title or '', size=30, color=th['primary'], bold=True, rtl=rtl,
            align=align, font=th['title_font'],
        )
        # accent underline
        line_y = y + Inches(0.72)
        line_w = Inches(1.4)
        line_x = (_W - _MARGIN - line_w) if rtl else _MARGIN
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, line_x, line_y, line_w, Inches(0.05))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(th['accent'])
        bar.line.fill.background()
        bar.shadow.inherit = False
        if subtitle:
            self._textbox(
                s, _MARGIN, line_y + Inches(0.15), _W - 2 * _MARGIN, Inches(0.4),
                subtitle, size=14, color=th['muted'], bold=False, rtl=rtl,
                align=align, font=th['body_font'],
            )

    def _fill_bullets(self, tf, bullets, th, rtl, *, size=18):
        align = _align(rtl)
        first = True
        for b in bullets or []:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            self._run(p, '• ' + b, size=size, color=th['text'], bold=False,
                      rtl=rtl, align=align, font=th['body_font'])
            p.space_after = Pt(8)
        if first:
            # empty frame — leave one blank para
            self._run(tf.paragraphs[0], '', size=size, color=th['text'], bold=False,
                      rtl=rtl, align=align, font=th['body_font'])

    def _picture(self, s, img_bytes, x, y, max_w, max_h):
        if not img_bytes:
            return
        try:
            with Image.open(io.BytesIO(img_bytes)) as im:
                iw, ih = im.size
            if iw <= 0 or ih <= 0:
                return
            ratio = min(float(max_w) / iw, float(max_h) / ih)
            w = int(iw * ratio)
            h = int(ih * ratio)
            s.shapes.add_picture(io.BytesIO(img_bytes), x, y, width=w, height=h)
        except Exception:
            pass  # soft-fail: bad image never fails the deck

    def _badge(self, s, th, x, y, size, text, rtl):
        shape = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(th['primary'])
        shape.line.fill.background()
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = str(text)
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = _rgb(th['bg'] if th['bg'] != th['primary'] else (255, 255, 255))
        run.font.name = th['title_font']

    # ── layouts ─────────────────────────────────────────────────────────

    def _title_font_size(self, title: str) -> int:
        """Scale hero title for long Persian lines."""
        n = len(title or '')
        if n > 90:
            return 28
        if n > 60:
            return 32
        if n > 40:
            return 36
        return 40

    def _lay_title(self, s, th, slide, rtl, img):
        title = slide.get('title') or ''
        subtitle = slide.get('subtitle') or ''
        size = self._title_font_size(title)

        if img:
            # Split hero: text column + framed image (RTL = text start-side).
            img_w, img_h = Inches(5.4), Inches(5.0)
            img_y = Inches(1.25)
            text_w = Inches(6.4)
            gap = Inches(0.35)
            if rtl:
                # Text on right (start), image on left
                img_x = _MARGIN
                text_x = _MARGIN + img_w + gap
            else:
                text_x = _MARGIN
                img_x = _W - _MARGIN - img_w
            # Surface frame behind image
            self._card(s, th, img_x - Inches(0.08), img_y - Inches(0.08),
                       img_w + Inches(0.16), img_h + Inches(0.16), accent_edge=False)
            self._picture(s, img, img_x, img_y, img_w, img_h)

            # Vertically center text block next to image
            text_top = Inches(2.0)
            box = s.shapes.add_textbox(text_x, text_top, text_w, Inches(2.4))
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._run(tf.paragraphs[0], title, size=size, color=th['primary'], bold=True,
                      rtl=rtl, align=_align(rtl), font=th['title_font'])
            # Accent bar under title block
            bar_w = Inches(1.6)
            bar_x = (text_x + text_w - bar_w) if rtl else text_x
            bar = s.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, bar_x, text_top + Inches(2.35), bar_w, Inches(0.06),
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = _rgb(th['accent'])
            bar.line.fill.background()
            bar.shadow.inherit = False
            if subtitle:
                self._textbox(
                    s, text_x, text_top + Inches(2.55), text_w, Inches(1.0),
                    subtitle, size=16, color=th['muted'], bold=False, rtl=rtl,
                    align=_align(rtl), font=th['body_font'],
                )
        else:
            # Centered title band
            band = s.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, 0, Inches(2.0), _W, Inches(3.4),
            )
            band.fill.solid()
            band.fill.fore_color.rgb = _rgb(th['surface'])
            band.line.fill.background()
            band.shadow.inherit = False
            tx, tw = Inches(1.0), Inches(11.3)
            box = s.shapes.add_textbox(tx, Inches(2.3), tw, Inches(1.8))
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._run(tf.paragraphs[0], title, size=size, color=th['primary'], bold=True,
                      rtl=rtl, align=PP_ALIGN.CENTER, font=th['title_font'])
            if subtitle:
                self._textbox(
                    s, tx, Inches(4.3), tw, Inches(0.8),
                    subtitle, size=16, color=th['muted'], bold=False, rtl=rtl,
                    align=PP_ALIGN.CENTER, font=th['body_font'],
                )

    def _lay_closing(self, s, th, slide, rtl, img):
        title = slide.get('title') or ''
        bullets = slide.get('bullets') or []
        self._textbox(
            s, Inches(1), Inches(2.3), Inches(11.3), Inches(1.4),
            title, size=40, color=th['primary'], bold=True, rtl=rtl,
            align=PP_ALIGN.CENTER, font=th['title_font'],
            vertical=MSO_ANCHOR.MIDDLE,
        )
        if bullets:
            box = s.shapes.add_textbox(Inches(2.5), Inches(4.0), Inches(8.3), Inches(2.2))
            tf = box.text_frame
            tf.word_wrap = True
            for i, b in enumerate(bullets[:5]):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                self._run(p, b, size=16, color=th['text'], bold=False, rtl=rtl,
                          align=PP_ALIGN.CENTER, font=th['body_font'])
                p.space_after = Pt(6)

    def _lay_section(self, s, th, slide, rtl, img):
        kicker = slide.get('kicker')
        title = slide.get('title') or ''
        # wide primary band
        band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.4), _W, Inches(2.6))
        band.fill.solid()
        band.fill.fore_color.rgb = _rgb(th['primary'])
        band.line.fill.background()
        band.shadow.inherit = False
        # contrast text on primary
        on_primary = th['bg'] if sum(th['primary']) < 400 else th['text']
        if kicker:
            self._textbox(
                s, _MARGIN, Inches(2.6), _W - 2 * _MARGIN, Inches(0.4),
                kicker, size=14, color=on_primary, bold=False, rtl=rtl,
                align=_align(rtl), font=th['body_font'],
            )
        self._textbox(
            s, _MARGIN, Inches(3.1), _W - 2 * _MARGIN, Inches(1.4),
            title, size=40, color=on_primary, bold=True, rtl=rtl,
            align=_align(rtl), font=th['title_font'],
        )

    def _lay_agenda(self, s, th, slide, rtl, img):
        self._title_block(s, th, slide.get('title') or '', rtl, kicker=slide.get('kicker'))
        items = slide.get('items') or slide.get('bullets') or []
        y = _BODY_TOP + Inches(0.15)
        for i, item in enumerate(items[:10]):
            self._badge(s, th, _MARGIN if not rtl else _W - _MARGIN - Inches(0.42),
                        y, Inches(0.42), str(i + 1), rtl)
            tx = _MARGIN + Inches(0.6) if not rtl else _MARGIN
            tw = _W - 2 * _MARGIN - Inches(0.7)
            self._textbox(
                s, tx, y, tw, Inches(0.45),
                item, size=18, color=th['text'], bold=False, rtl=rtl,
                align=_align(rtl), font=th['body_font'], vertical=MSO_ANCHOR.MIDDLE,
            )
            y += Inches(0.52)

    def _lay_content(self, s, th, slide, rtl, img):
        self._title_block(s, th, slide.get('title') or '', rtl,
                          kicker=slide.get('kicker'), subtitle=slide.get('subtitle'))
        bullets = slide.get('bullets') or []
        body_w = Inches(7.0) if img else (_W - 2 * _MARGIN)
        if img and rtl:
            bx = _MARGIN + Inches(5.3)
        else:
            bx = _MARGIN
        bbox = s.shapes.add_textbox(bx, _BODY_TOP + Inches(0.15), body_w, Inches(4.8))
        tf = bbox.text_frame
        tf.word_wrap = True
        self._fill_bullets(tf, bullets, th, rtl, size=18)
        if img:
            img_x = _MARGIN if rtl else Inches(8.2)
            self._picture(s, img, img_x, _BODY_TOP, _IMG_BOX, _IMG_BOX)

    def _lay_content_image(self, s, th, slide, rtl, img):
        self._lay_content(s, th, slide, rtl, img)

    def _side_has_content(self, side) -> bool:
        if not isinstance(side, dict):
            return False
        if (side.get('title') or '').strip():
            return True
        return bool(side.get('bullets'))

    def _resolve_two_columns(self, slide):
        """Return (left, right) sides with content, or None → fall back to content."""
        cols = slide.get('columns') if isinstance(slide.get('columns'), dict) else {}
        left = dict(cols.get('left') or {})
        right = dict(cols.get('right') or {})
        # Normalize alternate keys some models emit
        for side in (left, right):
            if not side.get('bullets') and side.get('items'):
                side['bullets'] = side['items']
            if not side.get('title') and side.get('label'):
                side['title'] = side['label']
            side['bullets'] = [str(b).strip() for b in (side.get('bullets') or []) if str(b).strip()]
            side['title'] = (side.get('title') or '').strip()

        if self._side_has_content(left) or self._side_has_content(right):
            # If one side empty, seed from leftover slide bullets
            bullets = [str(b).strip() for b in (slide.get('bullets') or []) if str(b).strip()]
            if not self._side_has_content(left) and bullets:
                mid = max(1, len(bullets) // 2)
                left = {'title': left.get('title') or '', 'bullets': bullets[:mid]}
                if not self._side_has_content(right):
                    right = {'title': right.get('title') or '', 'bullets': bullets[mid:]}
            return left, right

        # No columns — synthesize from bullets / features / steps
        bullets = [str(b).strip() for b in (slide.get('bullets') or []) if str(b).strip()]
        if not bullets and slide.get('features'):
            bullets = [
                ((f.get('title') or '') + (': ' + f['body'] if f.get('body') else '')).strip()
                for f in (slide.get('features') or []) if isinstance(f, dict)
            ]
            bullets = [b for b in bullets if b]
        if not bullets and slide.get('steps'):
            bullets = [
                ((f.get('title') or '') + (': ' + f['body'] if f.get('body') else '')).strip()
                for f in (slide.get('steps') or []) if isinstance(f, dict)
            ]
            bullets = [b for b in bullets if b]
        if len(bullets) >= 2:
            mid = (len(bullets) + 1) // 2
            return (
                {'title': '', 'bullets': bullets[:mid]},
                {'title': '', 'bullets': bullets[mid:]},
            )
        return None

    def _lay_two_column(self, s, th, slide, rtl, img):
        pair = self._resolve_two_columns(slide)
        if pair is None:
            # Empty two-column shell is worse than a normal content slide
            self._lay_content(s, th, slide, rtl, img)
            return
        left, right = pair
        self._title_block(s, th, slide.get('title') or '', rtl, kicker=slide.get('kicker'))
        # Body below title chrome
        body_y = _BODY_TOP + Inches(0.15)
        card_h = Inches(4.7)
        card_w = Inches(5.85)
        gap = Inches(0.35)
        # physical L/R; for RTL put "left" content on the right (reading start)
        if rtl:
            x_start, x_end = _MARGIN + card_w + gap, _MARGIN
            start, end = left, right
        else:
            x_start, x_end = _MARGIN, _MARGIN + card_w + gap
            start, end = left, right
        self._column_card(s, th, x_start, body_y, card_w, card_h, start, rtl)
        self._column_card(s, th, x_end, body_y, card_w, card_h, end, rtl)

    def _lay_comparison(self, s, th, slide, rtl, img):
        self._lay_two_column(s, th, slide, rtl, img)

    def _column_card(self, s, th, x, y, w, h, side, rtl):
        side = side or {}
        title = (side.get('title') or '').strip()
        bullets = [str(b).strip() for b in (side.get('bullets') or []) if str(b).strip()]
        # If only a title and no bullets, treat title as the body line
        if title and not bullets:
            bullets = [title]
            title = ''
        self._card(s, th, x, y, w, h)
        pad = Inches(0.28)
        cursor_y = y + Inches(0.22)
        if title:
            self._textbox(
                s, x + pad, cursor_y, w - 2 * pad, Inches(0.55),
                title, size=17, color=th['primary'], bold=True, rtl=rtl,
                align=_align(rtl), font=th['title_font'],
            )
            cursor_y += Inches(0.6)
            # thin accent under column title
            bar_w = Inches(0.9)
            bar_x = (x + w - pad - bar_w) if rtl else (x + pad)
            bar = s.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, bar_x, cursor_y - Inches(0.08), bar_w, Inches(0.04),
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = _rgb(th['accent'])
            bar.line.fill.background()
            bar.shadow.inherit = False
        body_h = h - (cursor_y - y) - Inches(0.25)
        box = s.shapes.add_textbox(x + pad, cursor_y, w - 2 * pad, max(body_h, Inches(1.0)))
        tf = box.text_frame
        tf.word_wrap = True
        if bullets:
            self._fill_bullets(tf, bullets[:8], th, rtl, size=14)
        else:
            # Never leave a blank card — soft placeholder (should be rare after resolve)
            self._run(tf.paragraphs[0], '—', size=14, color=th['muted'], bold=False,
                      rtl=rtl, align=_align(rtl), font=th['body_font'])

    def _lay_metrics(self, s, th, slide, rtl, img):
        self._title_block(s, th, slide.get('title') or '', rtl, kicker=slide.get('kicker'))
        metrics = slide.get('metrics') or []
        if not metrics:
            # fallback to bullets as pseudo-metrics
            for b in (slide.get('bullets') or [])[:4]:
                metrics.append({'value': b[:20], 'label': '', 'delta': None})
        n = max(len(metrics), 1)
        n = min(n, 4)
        gap = Inches(0.25)
        total_w = _W - 2 * _MARGIN
        card_w = (total_w - gap * (n - 1)) / n
        y = _BODY_TOP + Inches(0.4)
        h = Inches(3.6)
        for i, m in enumerate(metrics[:n]):
            x = _MARGIN + i * (card_w + gap)
            self._card(s, th, x, y, card_w, h)
            # value — prefer LTR so numbers don't reverse
            self._textbox(
                s, x + Inches(0.2), y + Inches(0.7), card_w - Inches(0.4), Inches(1.1),
                m.get('value') or '', size=36, color=th['primary'], bold=True,
                rtl=False, align=PP_ALIGN.CENTER, font=th['title_font'],
                vertical=MSO_ANCHOR.MIDDLE,
            )
            self._textbox(
                s, x + Inches(0.2), y + Inches(1.9), card_w - Inches(0.4), Inches(0.7),
                m.get('label') or '', size=14, color=th['text'], bold=False,
                rtl=rtl, align=PP_ALIGN.CENTER, font=th['body_font'],
            )
            if m.get('delta'):
                self._textbox(
                    s, x + Inches(0.2), y + Inches(2.7), card_w - Inches(0.4), Inches(0.45),
                    m['delta'], size=13, color=th['accent'], bold=True,
                    rtl=False, align=PP_ALIGN.CENTER, font=th['body_font'],
                )

    def _lay_process(self, s, th, slide, rtl, img):
        self._title_block(s, th, slide.get('title') or '', rtl, kicker=slide.get('kicker'))
        steps = slide.get('steps') or []
        if not steps and slide.get('bullets'):
            steps = [{'title': b, 'body': '', 'icon': None} for b in slide['bullets'][:5]]
        n = min(len(steps), 5) or 1
        gap = Inches(0.2)
        total_w = _W - 2 * _MARGIN
        card_w = (total_w - gap * (n - 1)) / n
        y = _BODY_TOP + Inches(0.25)
        h = Inches(4.6)
        order = list(range(n))
        if rtl:
            order = list(reversed(order))
        for visual_i, step_i in enumerate(order):
            step = steps[step_i]
            x = _MARGIN + visual_i * (card_w + gap)
            self._card(s, th, x, y, card_w, h)
            self._badge(s, th, x + (card_w - Inches(0.45)) / 2, y + Inches(0.3),
                        Inches(0.45), str(step_i + 1), rtl)
            self._textbox(
                s, x + Inches(0.15), y + Inches(1.0), card_w - Inches(0.3), Inches(0.9),
                step.get('title') or '', size=15, color=th['primary'], bold=True,
                rtl=rtl, align=PP_ALIGN.CENTER, font=th['title_font'],
            )
            self._textbox(
                s, x + Inches(0.15), y + Inches(2.0), card_w - Inches(0.3), Inches(2.2),
                step.get('body') or '', size=12, color=th['text'], bold=False,
                rtl=rtl, align=PP_ALIGN.CENTER, font=th['body_font'],
            )

    def _lay_timeline(self, s, th, slide, rtl, img):
        self._title_block(s, th, slide.get('title') or '', rtl, kicker=slide.get('kicker'))
        events = slide.get('events') or []
        if not events and slide.get('bullets'):
            events = [{'label': b, 'detail': '', 'when': None} for b in slide['bullets'][:6]]
        n = min(len(events), 6) or 1
        # horizontal axis
        axis_y = Inches(3.6)
        axis = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, _MARGIN, axis_y, _W - 2 * _MARGIN, Inches(0.06),
        )
        axis.fill.solid()
        axis.fill.fore_color.rgb = _rgb(th['primary'])
        axis.line.fill.background()
        axis.shadow.inherit = False

        slot_w = (_W - 2 * _MARGIN) / n
        for i, ev in enumerate(events[:n]):
            # visual order: RTL starts from right
            vi = (n - 1 - i) if rtl else i
            cx = _MARGIN + vi * slot_w + slot_w / 2
            # node
            node_s = Inches(0.28)
            node = s.shapes.add_shape(
                MSO_SHAPE.OVAL, cx - node_s / 2, axis_y - Inches(0.11), node_s, node_s,
            )
            node.fill.solid()
            node.fill.fore_color.rgb = _rgb(th['accent'])
            node.line.fill.background()
            node.shadow.inherit = False

            above = (i % 2 == 0)
            ty = axis_y - Inches(1.7) if above else axis_y + Inches(0.4)
            when = ev.get('when') or ''
            label = ev.get('label') or ''
            detail = ev.get('detail') or ''
            block = (when + '\n' if when else '') + label + (('\n' + detail) if detail else '')
            self._textbox(
                s, cx - slot_w / 2 + Inches(0.05), ty, slot_w - Inches(0.1), Inches(1.5),
                block, size=11, color=th['text'], bold=False, rtl=rtl,
                align=PP_ALIGN.CENTER, font=th['body_font'],
            )

    def _lay_quote(self, s, th, slide, rtl, img):
        quote = slide.get('quote') or slide.get('title') or ''
        attr = slide.get('attribution') or ''
        # large surface card
        self._card(s, th, Inches(1.2), Inches(1.6), Inches(10.9), Inches(4.4),
                   accent_edge=True)
        # decorative mark
        mark = '«' if rtl else '"'
        self._textbox(
            s, Inches(1.5), Inches(1.8), Inches(1.2), Inches(0.8),
            mark, size=48, color=th['accent'], bold=True, rtl=False,
            align=PP_ALIGN.LEFT, font=th['title_font'],
        )
        self._textbox(
            s, Inches(1.8), Inches(2.6), Inches(9.5), Inches(2.2),
            quote, size=26, color=th['text'], bold=False, rtl=rtl,
            align=PP_ALIGN.CENTER, font=th['title_font'],
            vertical=MSO_ANCHOR.MIDDLE,
        )
        if attr:
            self._textbox(
                s, Inches(1.8), Inches(5.0), Inches(9.5), Inches(0.5),
                '— ' + attr, size=14, color=th['muted'], bold=False, rtl=rtl,
                align=PP_ALIGN.CENTER, font=th['body_font'],
            )

    def _lay_feature_grid(self, s, th, slide, rtl, img):
        self._title_block(s, th, slide.get('title') or '', rtl, kicker=slide.get('kicker'))
        features = slide.get('features') or slide.get('steps') or []
        if not features and slide.get('bullets'):
            features = [{'title': b, 'body': '', 'icon': None} for b in slide['bullets'][:6]]
        n = min(len(features), 6)
        cols = 3 if n > 4 else (2 if n > 1 else 1)
        rows = (n + cols - 1) // cols
        gap = Inches(0.22)
        total_w = _W - 2 * _MARGIN
        total_h = Inches(4.7)
        card_w = (total_w - gap * (cols - 1)) / cols
        card_h = (total_h - gap * (rows - 1)) / max(rows, 1)
        for i, feat in enumerate(features[:n]):
            r, c = divmod(i, cols)
            # RTL: mirror column index
            vc = (cols - 1 - c) if rtl else c
            x = _MARGIN + vc * (card_w + gap)
            y = _BODY_TOP + Inches(0.1) + r * (card_h + gap)
            self._card(s, th, x, y, card_w, card_h)
            # icon badge
            self._badge(s, th, x + Inches(0.2), y + Inches(0.2), Inches(0.36),
                        str(i + 1), rtl)
            self._textbox(
                s, x + Inches(0.65), y + Inches(0.2), card_w - Inches(0.85), Inches(0.45),
                feat.get('title') or '', size=15, color=th['primary'], bold=True,
                rtl=rtl, align=_align(rtl), font=th['title_font'],
                vertical=MSO_ANCHOR.MIDDLE,
            )
            self._textbox(
                s, x + Inches(0.2), y + Inches(0.75), card_w - Inches(0.4),
                card_h - Inches(0.95),
                feat.get('body') or '', size=12, color=th['text'], bold=False,
                rtl=rtl, align=_align(rtl), font=th['body_font'],
            )

    def _lay_image_hero(self, s, th, slide, rtl, img):
        title = slide.get('title') or ''
        caption = slide.get('caption') or slide.get('subtitle') or ''
        # dark-ish overlay band at bottom for title
        if img:
            self._picture(s, img, Inches(0.4), Inches(0.35), Inches(12.5), Inches(5.5))
        band = s.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, Inches(5.5), _W, Inches(2.0),
        )
        band.fill.solid()
        band.fill.fore_color.rgb = _rgb(th['primary'])
        band.line.fill.background()
        band.shadow.inherit = False
        on = th['bg'] if sum(th['primary']) < 400 else th['text']
        self._textbox(
            s, _MARGIN, Inches(5.7), _W - 2 * _MARGIN, Inches(0.8),
            title, size=28, color=on, bold=True, rtl=rtl,
            align=_align(rtl), font=th['title_font'],
        )
        if caption:
            self._textbox(
                s, _MARGIN, Inches(6.5), _W - 2 * _MARGIN, Inches(0.5),
                caption, size=14, color=on, bold=False, rtl=rtl,
                align=_align(rtl), font=th['body_font'],
            )

    def _lay_chart(self, s, th, slide, rtl, img):
        self._title_block(s, th, slide.get('title') or '', rtl, kicker=slide.get('kicker'))
        chart = slide.get('chart')
        if not chart:
            # degrade to metrics/content
            if slide.get('metrics'):
                self._lay_metrics(s, th, slide, rtl, img)
            else:
                self._lay_content(s, th, slide, rtl, img)
            return
        try:
            self._add_chart(s, th, chart)
        except Exception:
            # never fail deck on chart
            metrics = []
            cats = chart.get('categories') or []
            vals = (chart.get('series') or [{}])[0].get('values') or []
            for c, v in zip(cats[:4], vals[:4]):
                metrics.append({'value': str(v), 'label': c, 'delta': None})
            slide = {**slide, 'metrics': metrics}
            self._lay_metrics(s, th, slide, rtl, img)

    def _add_chart(self, s, th, chart):
        data = CategoryChartData()
        data.categories = chart['categories']
        for series in chart['series']:
            data.add_series(series['name'], series['values'])
        ctype = {
            'bar': XL_CHART_TYPE.COLUMN_CLUSTERED,
            'line': XL_CHART_TYPE.LINE_MARKERS,
            'pie': XL_CHART_TYPE.PIE,
        }.get(chart.get('type') or 'bar', XL_CHART_TYPE.COLUMN_CLUSTERED)
        x, y, w, h = _MARGIN, _BODY_TOP + Inches(0.1), _W - 2 * _MARGIN, Inches(4.8)
        graphic = s.shapes.add_chart(ctype, x, y, w, h, data)
        ch = graphic.chart
        if chart.get('type') != 'pie':
            try:
                ch.has_legend = len(chart.get('series') or []) > 1
                if ch.has_legend:
                    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
                    ch.legend.include_in_layout = False
            except Exception:
                pass
        # try color first series
        try:
            plot = ch.plots[0]
            if hasattr(plot, 'series') and plot.series:
                ser = plot.series[0]
                if chart.get('type') != 'pie':
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = _rgb(th['primary'])
        except Exception:
            pass
